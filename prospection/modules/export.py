"""
Module Export — Génération du fichier Excel professionnel en 3 onglets.
"""

import sys
import os
from datetime import datetime
from typing import List

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from config import BASE_DIR, creer_logger, normaliser_nom_fichier
from modules.modeles import Parapharmacie

logger = creer_logger("export")

# Palette de couleurs
COULEUR_ENTETE = "1F3864"
COULEUR_LIGNES_PAIRES = "F5F5F5"
COULEUR_PROSPECTION = "FFFACD"
COULEUR_MATCH_FAIBLE = "FF9999"   # rouge clair : < 70%
COULEUR_MATCH_MOYEN = "FFD580"    # orange clair : 70–90%
COULEUR_MATCH_BON = "90EE90"      # vert clair   : > 90%

# Définition des colonnes de l'onglet principal
COLONNES = [
    ("A", "Nom"),
    ("B", "Adresse"),
    ("C", "Ville"),
    ("D", "Code Postal"),
    ("E", "Téléphone"),
    ("F", "Email"),
    ("G", "Site Web"),
    ("H", "SIRET"),
    ("I", "CA (€)"),
    ("J", "Année Bilan"),
    ("K", "Dirigeant"),
    ("L", "Confiance Match (%)"),
    ("M", "Source Données"),
    ("N", "Statut Prospection"),
    ("O", "Notes"),
    ("P", "Date Contact"),
]

# Colonnes à remplir manuellement → fond jaune
COLONNES_PROSPECTION = {"N", "O", "P"}


def _valeur_colonne(p: Parapharmacie, lettre: str):
    """Retourne la valeur d'un objet Parapharmacie pour une lettre de colonne."""
    mapping = {
        "A": p.nom,
        "B": p.adresse,
        "C": p.ville,
        "D": p.code_postal,
        "E": p.telephone,
        "F": p.email,
        "G": p.site_web,
        "H": p.siret,
        "I": p.chiffre_affaires,
        "J": p.annee_bilan,
        "K": p.dirigeant,
        "L": p.confiance_match if p.confiance_match else None,
        "M": p.source_principale,
        "N": p.statut_prospection,
        "O": p.notes,
        "P": p.date_contact,
    }
    val = mapping.get(lettre)
    return val if val is not None else ""


def creer_onglet_parapharmacies(
    wb: openpyxl.Workbook, donnees: List[Parapharmacie]
) -> None:
    """Crée et remplit l'onglet principal avec formatage professionnel."""
    ws = wb.active
    ws.title = "Parapharmacies"

    fill_entete = PatternFill("solid", fgColor=COULEUR_ENTETE)
    font_entete = Font(color="FFFFFF", bold=True)
    fill_paire = PatternFill("solid", fgColor=COULEUR_LIGNES_PAIRES)
    fill_prospection = PatternFill("solid", fgColor=COULEUR_PROSPECTION)

    # En-têtes
    for col_letter, col_nom in COLONNES:
        cell = ws[f"{col_letter}1"]
        cell.value = col_nom
        cell.fill = fill_entete
        cell.font = font_entete
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Données
    for row_idx, p in enumerate(donnees, start=2):
        for col_letter, _ in COLONNES:
            cell = ws[f"{col_letter}{row_idx}"]
            cell.value = _valeur_colonne(p, col_letter)

            # Lignes alternées
            if row_idx % 2 == 0:
                cell.fill = fill_paire

            # Colonnes prospection en jaune (écrase le gris des lignes paires)
            if col_letter in COLONNES_PROSPECTION:
                cell.fill = fill_prospection

        # Formatage conditionnel colonne L (score de confiance match)
        cell_l = ws[f"L{row_idx}"]
        try:
            val = int(cell_l.value) if cell_l.value else 0
            if val > 0:
                if val < 70:
                    cell_l.fill = PatternFill("solid", fgColor=COULEUR_MATCH_FAIBLE)
                elif val <= 90:
                    cell_l.fill = PatternFill("solid", fgColor=COULEUR_MATCH_MOYEN)
                else:
                    cell_l.fill = PatternFill("solid", fgColor=COULEUR_MATCH_BON)
        except (ValueError, TypeError):
            pass

    # Ajustement automatique des largeurs de colonnes
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 50)

    # Filtre automatique sur toutes les colonnes
    ws.auto_filter.ref = ws.dimensions

    # Ligne d'en-tête figée
    ws.freeze_panes = "A2"


def creer_onglet_resume(
    wb: openpyxl.Workbook,
    donnees: List[Parapharmacie],
    zone: str,
) -> None:
    """Crée l'onglet de résumé avec statistiques."""
    ws = wb.create_sheet("Résumé")

    total = len(donnees)
    avec_siret = sum(1 for p in donnees if p.siret)
    avec_ca = sum(1 for p in donnees if p.chiffre_affaires)
    avec_email = sum(1 for p in donnees if p.email)

    fill_titre = PatternFill("solid", fgColor=COULEUR_ENTETE)
    font_titre = Font(color="FFFFFF", bold=True, size=13)
    font_label = Font(bold=True)
    font_valeur = Font()

    # Titre
    ws["A1"] = "Rapport de prospection"
    ws["A1"].fill = fill_titre
    ws["A1"].font = font_titre
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("A1:B1")
    ws.row_dimensions[1].height = 24

    # Données du résumé
    pct = lambda n: f"{n/total*100:.1f}%" if total else "0%"
    lignes = [
        ("Zone recherchée", zone),
        ("Date génération", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Total trouvées", str(total)),
        ("Avec SIRET", f"{avec_siret} ({pct(avec_siret)})"),
        ("Avec CA", f"{avec_ca} ({pct(avec_ca)})"),
        ("Avec email", f"{avec_email} ({pct(avec_email)})"),
    ]

    for i, (label, valeur) in enumerate(lignes, start=2):
        ws[f"A{i}"] = label
        ws[f"A{i}"].font = font_label
        ws[f"B{i}"] = valeur
        ws[f"B{i}"].font = font_valeur
        ws.row_dimensions[i].height = 18

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 30
    ws.freeze_panes = "A2"


def creer_onglet_sources(wb: openpyxl.Workbook) -> None:
    """Crée l'onglet d'information sur les sources de données."""
    ws = wb.create_sheet("Sources et Limites")

    fill_titre = PatternFill("solid", fgColor=COULEUR_ENTETE)
    font_titre = Font(color="FFFFFF", bold=True, size=12)
    font_source = Font(bold=True)

    ws["A1"] = "Sources de données et limites connues"
    ws["A1"].fill = fill_titre
    ws["A1"].font = font_titre
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("A1:B1")
    ws.row_dimensions[1].height = 22

    sources = [
        (
            "OpenStreetMap",
            "Base principale, données collaboratives. "
            "Peut être incomplet pour les petites villes.",
        ),
        (
            "INSEE Sirene",
            "Données officielles. Fiables. "
            "Nécessite une clé API gratuite sur api.insee.fr",
        ),
        (
            "Pappers",
            "Chiffre d'affaires avec 1 à 2 ans de retard possible. "
            "Absent si l'entreprise n'a pas déposé son bilan.",
        ),
        (
            "Emails",
            "Trouvés uniquement si visibles publiquement sur le site. "
            "Absent pour la majorité des petites structures.",
        ),
    ]

    for i, (source, description) in enumerate(sources, start=2):
        ws[f"A{i}"] = source
        ws[f"A{i}"].font = font_source
        ws[f"B{i}"] = description
        ws[f"B{i}"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 45

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 72


def generer(
    donnees: List[Parapharmacie],
    zone: str,
    partiel: bool = False,
) -> str:
    """
    Génère le fichier Excel complet.
    Si partiel=True, le nom indique que les données sont incomplètes.
    Retourne le chemin absolu du fichier généré.
    """
    os.makedirs(os.path.join(BASE_DIR, "exports"), exist_ok=True)

    suffixe = "_partiel" if partiel else ""
    nom_fichier = os.path.join(
        BASE_DIR, "exports",
        f"prospection_{normaliser_nom_fichier(zone)}_"
        f"{datetime.now().strftime('%Y%m%d')}{suffixe}.xlsx"
    )

    wb = openpyxl.Workbook()
    creer_onglet_parapharmacies(wb, donnees)
    creer_onglet_resume(wb, donnees, zone)
    creer_onglet_sources(wb)

    wb.save(nom_fichier)
    logger.info(f"Excel généré : {nom_fichier} ({len(donnees)} entrées)")

    return nom_fichier


if __name__ == "__main__":
    import glob
    import json

    print("=" * 55)
    print("  TEST MODULE EXPORT — Lyon")
    print("=" * 55)

    # Charger les données les plus enrichies disponibles
    for pattern in [
        os.path.join(BASE_DIR, "cache", "progression_lyon_phase3.json"),
        os.path.join(BASE_DIR, "cache", "progression_lyon_phase2.json"),
    ]:
        fichiers = glob.glob(pattern)
        if fichiers:
            break
    else:
        fichiers = glob.glob(os.path.join(BASE_DIR, "cache", "osm_lyon_*.json"))

    if fichiers:
        fichier = sorted(fichiers)[-1]
        print(f"  [INFO] Chargement : {os.path.basename(fichier)}")
        with open(fichier, 'r', encoding='utf-8') as f:
            donnees = [Parapharmacie(**d) for d in json.load(f)]
    else:
        # Données de démonstration si aucun cache
        print("  [INFO] Données de démonstration")
        donnees = [
            Parapharmacie(
                nom="Carré Santé Beauté Monplaisir",
                adresse="", ville="Lyon", code_postal="69008",
                telephone="04 78 78 80 78",
                email="monplaisir@carre-sante.com",
                siret="12345678900001",
                nom_legal="CARRE SANTE BEAUTE SARL",
                dirigeant="Sophie Martin",
                chiffre_affaires="480000", annee_bilan="2022",
                source_ca="Pappers",
                confiance_match=92, statut_match="AUTO",
                source_principale="OpenStreetMap",
            ),
            Parapharmacie(
                nom="Parapharmacie Valmy",
                adresse="", ville="Lyon", code_postal="69009",
                telephone="",
                confiance_match=65, statut_match="A_VERIFIER",
                source_principale="OpenStreetMap",
            ),
            Parapharmacie(
                nom="Grande Droguerie Lyonnaise",
                adresse="", ville="Lyon", code_postal="69002",
                telephone="04 78 42 78 52",
                email="droguerielyonnaise@orange.fr",
                site_web="https://www.la-droguerie.com/",
                confiance_match=0, statut_match="NON_TROUVE",
                source_principale="OpenStreetMap",
            ),
        ]

    chemin = generer(donnees, "Lyon")
    print(f"\n  [OK] Fichier Excel généré : {os.path.basename(chemin)}")
    print(f"  → 3 onglets : Parapharmacies | Résumé | Sources et Limites")
    print(f"  → {len(donnees)} entrées exportées")
    print(f"  → Chemin : {chemin}")
