"""
Import du fichier Excel des pharmacies : Copie_de_Pharma_pour_vadim.xlsx
Gère le mapping des colonnes, la détection des pharmacies en cible et l'anti-doublon.
"""

import os
import logging
from typing import Callable, Optional

import openpyxl

from database import (
    inserer_pharmacie,
    TRANCHES_EN_CIBLE,
    CA_VALEURS_ESTIMEES,
    pharmacie_existe,
)

logger = logging.getLogger(__name__)


COLONNES_MAPPING = {
    "PHARMACIE": "nom",
    "ADRESSE": "adresse",
    "VILLE": "ville",
    "CODEPOSTAL": "code_postal",
    "dpt": "departement",
    "TELEPHONE": "telephone",
    "CA_TTC_D_C": "ca_tranche",
    "classement": "classement",
}


def normaliser_ca(valeur) -> str:
    """Normalise la valeur de la tranche CA en chaîne propre."""
    if valeur is None:
        return ""
    s = str(valeur).strip().upper()
    return s


def est_en_cible(ca_tranche: str) -> bool:
    """Retourne True si la tranche de CA correspond à une pharmacie en cible."""
    return ca_tranche.upper() in TRANCHES_EN_CIBLE


def importer_excel(
    db_path: str,
    chemin_excel: str,
    callback_progress: Optional[Callable[[int, int, str], None]] = None,
) -> dict:
    """
    Importe le fichier Excel dans la base SQLite.
    callback_progress(actuel, total, message) est appelé régulièrement pour mettre à jour l'UI.
    Retourne un dict avec les compteurs : insérées, doublons, erreurs.
    """
    if not os.path.exists(chemin_excel):
        raise FileNotFoundError(f"Fichier introuvable : {chemin_excel}")

    wb = openpyxl.load_workbook(chemin_excel, read_only=True, data_only=True)
    ws = wb.active

    # Lecture de la ligne d'en-tête
    headers = {}
    for cell in next(ws.iter_rows(min_row=1, max_row=1)):
        if cell.value is not None:
            headers[str(cell.value).strip()] = cell.column - 1  # index 0-based

    # Toutes les lignes de données
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    total = len(rows)

    compteurs = {"inserees": 0, "doublons": 0, "erreurs": 0}
    nom = ""  # Initialisé pour le callback de progression

    for i, row in enumerate(rows):
        try:
            def col(nom_colonne: str, defaut=""):
                idx = headers.get(nom_colonne)
                if idx is None:
                    return defaut
                val = row[idx]
                return str(val).strip() if val is not None else defaut

            nom = col("PHARMACIE")
            if not nom:
                continue

            ville = col("VILLE")
            code_postal = col("CODEPOSTAL")

            # Vérification doublon
            if pharmacie_existe(db_path, nom, ville, code_postal):
                compteurs["doublons"] += 1
            else:
                ca_tranche = normaliser_ca(col("CA_TTC_D_C"))
                en_cible = 1 if est_en_cible(ca_tranche) else 0
                ca_value = CA_VALEURS_ESTIMEES.get(ca_tranche.upper())

                data = {
                    "nom": nom,
                    "adresse": col("ADRESSE") or None,
                    "ville": ville or None,
                    "code_postal": code_postal or None,
                    "departement": col("dpt") or None,
                    "telephone": col("TELEPHONE") or None,
                    "ca_tranche": ca_tranche or None,
                    "ca_value": ca_value,
                    "en_cible": en_cible,
                    "source": "import_excel",
                    "fax": None,
                    "email": None,
                    "nom_decidant": None,
                    "prenom_decidant": None,
                    "telephone_decidant": None,
                    "email_decidant": None,
                    "date_dernier_contact": None,
                }

                result = inserer_pharmacie(db_path, data)
                if result:
                    compteurs["inserees"] += 1
                else:
                    compteurs["doublons"] += 1

        except Exception as e:
            compteurs["erreurs"] += 1
            logger.warning(f"Erreur ligne {i + 2} : {e}")

        if callback_progress and (i % 50 == 0 or i == total - 1):
            msg = f"Traitement {i + 1}/{total} — {nom if nom else '...'}"
            callback_progress(i + 1, total, msg)

    wb.close()
    logger.info(
        f"Import terminé : {compteurs['inserees']} insérées, "
        f"{compteurs['doublons']} doublons, {compteurs['erreurs']} erreurs"
    )
    return compteurs
